"""Build the professional hypertension research report (Word .docx and PDF).

The script writes ``Hypertension_Research_Report.docx`` with python-docx and then
converts it to ``Hypertension_Research_Report.pdf`` with LibreOffice, so that both
deliverables share exactly the same layout.

Usage:
    python build_report.py            # figures must already exist in ./figures
"""

import subprocess
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).parent
FIGURES = BASE / "figures"
DOCX = BASE / "Hypertension_Research_Report.docx"

NAVY = RGBColor(0x12, 0x36, 0x5B)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
GREY = RGBColor(0x55, 0x5F, 0x68)
HEADER_FILL = "12365B"
BAND_FILL = "EDF2F6"


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, *, bold=False, colour=None, size=9.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if colour is not None:
        run.font.color.rgb = colour


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, header, bold=True, colour=RGBColor(0xFF, 0xFF, 0xFF), size=9.5)
        shade(cell, HEADER_FILL)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value, bold=(index == 0))
            if row_index % 2 == 1:
                shade(cells[index], BAND_FILL)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    # keep the header row visible when a table spans pages
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    return table


def add_figure(document, filename, caption, width_cm=16.0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(FIGURES / filename), width=Cm(width_cm))
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def add_horizontal_rule(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), HEADER_FILL)
    borders.append(bottom)
    p_pr.append(borders)


# --------------------------------------------------------------------------- #
# document scaffolding
# --------------------------------------------------------------------------- #
def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, colour, space_before in (
        ("Title", 30, NAVY, 0),
        ("Heading 1", 16, NAVY, 16),
        ("Heading 2", 13, BLUE, 12),
        ("Heading 3", 11.5, BLUE, 10),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = colour
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(6)

    caption = document.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = GREY

    quote = document.styles["Intense Quote"]
    quote.font.size = Pt(10.5)
    quote.font.color.rgb = NAVY


def configure_page(document):
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    return section


def add_header_footer(section, title):
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(title)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_horizontal_rule(paragraph)

    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.text = ""
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = foot.add_run("Page ")
    label.font.size = Pt(8.5)
    label.font.color.rgb = GREY
    add_field(foot, " PAGE ")
    of_run = foot.add_run(" of ")
    of_run.font.size = Pt(8.5)
    of_run.font.color.rgb = GREY
    add_field(foot, " NUMPAGES ")
    for run in foot.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = GREY


def add_cover(document):
    for _ in range(3):
        document.add_paragraph()

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("CLINICAL RESEARCH REPORT")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = BLUE

    title = document.add_paragraph("Hypertension", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "Epidemiology, Pathophysiology, Diagnosis and Evidence-Based Management"
    )
    run.font.size = Pt(14)
    run.font.color.rgb = GREY

    rule = document.add_paragraph()
    add_horizontal_rule(rule)

    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(FIGURES / "fig2_care_cascade.png"), width=Cm(15.0))

    for _ in range(2):
        document.add_paragraph()

    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Prepared as an evidence review of the WHO Global Report on Hypertension,\n"
        "the 2025 ACC/AHA High Blood Pressure Guideline and the 2024 ESC Guidelines\n"
        "for the Management of Elevated Blood Pressure and Hypertension"
    )
    run.font.size = Pt(10.5)
    run.font.color.rgb = GREY

    document.add_page_break()


def add_table_of_contents(document):
    document.add_heading("Table of Contents", level=1)
    entries = [
        "1. Executive Summary",
        "2. Introduction and Definitions",
        "3. Global Burden and Epidemiology",
        "4. Pathophysiology and Risk Factors",
        "5. Clinical Presentation and Complications",
        "6. Diagnosis and Blood-Pressure Measurement",
        "7. Non-Pharmacological Management",
        "8. Pharmacological Management",
        "9. Special Populations and Resistant Hypertension",
        "10. Hypertensive Crises",
        "11. Prevention and Public-Health Strategy",
        "12. Conclusion",
        "13. References",
    ]
    for entry in entries:
        paragraph = document.add_paragraph(entry)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.runs[0].font.size = Pt(11)
    document.add_page_break()


# --------------------------------------------------------------------------- #
# content
# --------------------------------------------------------------------------- #
def add_body(document):
    document.add_heading("1. Executive Summary", level=1)
    document.add_paragraph(
        "Hypertension, or persistently raised arterial blood pressure, is the single largest "
        "modifiable contributor to cardiovascular disease and premature death worldwide. The "
        "World Health Organization estimates that about 1.28 billion adults aged 30-79 years "
        "live with hypertension, roughly one in three adults in that age band, and that two "
        "thirds of them live in low- and middle-income countries. The condition is usually "
        "silent: close to half of affected adults are unaware of their diagnosis, only about "
        "42 per cent receive treatment, and only about 21 per cent achieve control."
    )
    document.add_paragraph(
        "The clinical case for action is strong and quantifiable. Meta-analyses of randomised "
        "trials show that lowering systolic blood pressure by 10 mm Hg reduces major "
        "cardiovascular events by roughly one fifth, stroke by about one quarter and heart "
        "failure by about one quarter to one third. Contemporary guidelines therefore emphasise "
        "accurate measurement, earlier risk-based treatment, lower treatment targets and "
        "simplified drug regimens built around single-pill combinations."
    )

    summary = document.add_paragraph(
        "Key message: hypertension is common, largely asymptomatic, inexpensive to detect and "
        "highly treatable. Most of the lost benefit is not caused by a lack of effective drugs "
        "but by gaps in diagnosis, treatment initiation, titration and long-term adherence.",
        style="Intense Quote",
    )
    summary.alignment = WD_ALIGN_PARAGRAPH.LEFT

    document.add_heading("2. Introduction and Definitions", level=1)
    document.add_paragraph(
        "Blood pressure is the force exerted by circulating blood on the arterial wall. It is "
        "reported as systolic pressure (during ventricular contraction) over diastolic pressure "
        "(during ventricular relaxation) in millimetres of mercury. Risk rises continuously from "
        "systolic values as low as 115 mm Hg, so any diagnostic threshold is a pragmatic "
        "convention chosen to identify people who benefit from intervention rather than a "
        "biological cut-off."
    )
    document.add_paragraph(
        "The American College of Cardiology and American Heart Association define hypertension "
        "as an office blood pressure of 130/80 mm Hg or higher; the 2025 guideline retains the "
        "categories introduced in 2017 but places greater weight on estimated cardiovascular "
        "risk, calculated with the PREVENT equations, when deciding who should start drug "
        "therapy. The 2024 European Society of Cardiology guideline keeps the traditional "
        "140/90 mm Hg diagnostic threshold but adds an intermediate category of 'elevated blood "
        "pressure' (120-139/70-89 mm Hg) in which treatment is decided by overall risk."
    )
    add_figure(
        document,
        "fig1_bp_categories.png",
        "Figure 1. Office blood-pressure categories in the ACC/AHA and ESC frameworks. "
        "Both systems now treat blood pressure as a continuous risk factor rather than a "
        "simple normal/abnormal dichotomy.",
    )

    add_table(
        document,
        ["Category", "Systolic (mm Hg)", "Diastolic (mm Hg)", "Usual initial action"],
        [
            ["Normal", "< 120", "and < 80", "Reassess periodically; reinforce healthy lifestyle"],
            ["Elevated", "120-129", "and < 80", "Structured lifestyle therapy; repeat measurement"],
            [
                "Stage 1 hypertension",
                "130-139",
                "or 80-89",
                "Lifestyle therapy; add drugs if cardiovascular risk is high",
            ],
            [
                "Stage 2 hypertension",
                "≥ 140",
                "or ≥ 90",
                "Lifestyle therapy plus drug treatment, usually two agents",
            ],
            [
                "Hypertensive crisis",
                "> 180",
                "and/or > 120",
                "Immediate reassessment; emergency care if organ damage is present",
            ],
        ],
        widths=[4.0, 3.0, 3.0, 6.0],
    )
    caption = document.add_paragraph(
        "Table 1. ACC/AHA blood-pressure categories for adults and the corresponding "
        "first management step.",
        style="Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("3. Global Burden and Epidemiology", level=1)
    document.add_paragraph(
        "Global prevalence of hypertension in adults aged 30-79 years doubled in absolute terms "
        "between 1990 and 2019, from roughly 650 million to 1.28 billion people, driven mainly "
        "by population growth, ageing and rising obesity. Age-standardised prevalence has been "
        "broadly stable at about one third of adults, but the geography has shifted decisively: "
        "prevalence has fallen in many high-income countries while rising in low- and "
        "middle-income regions, which now carry roughly two thirds of the global caseload."
    )
    document.add_paragraph(
        "Raised blood pressure is implicated in around 10 million deaths each year, including "
        "about half of all deaths from stroke and ischaemic heart disease. The care cascade "
        "shown in Figure 2 explains most of that toll: attrition occurs at every step from "
        "awareness to sustained control. Modelling by WHO suggests that scaling effective "
        "treatment coverage to 50 per cent of adults with hypertension could prevent about 76 "
        "million deaths between 2023 and 2050."
    )
    add_figure(
        document,
        "fig2_care_cascade.png",
        "Figure 2. The hypertension care cascade. Roughly half of affected adults are unaware "
        "of their condition and only about one in five reaches blood-pressure control.",
    )

    add_table(
        document,
        ["Indicator", "Global estimate", "Practical implication"],
        [
            ["Adults living with hypertension (30-79 y)", "~1.28 billion", "One in three adults in this age group"],
            ["Living in low- and middle-income countries", "~2 of every 3", "Health-system capacity is the binding constraint"],
            ["Aware of their diagnosis", "~54%", "Opportunistic screening remains essential"],
            ["Receiving antihypertensive treatment", "~42%", "Treatment initiation is delayed or never happens"],
            ["Blood pressure controlled on treatment", "~21%", "Under-titration and non-adherence dominate"],
            ["Annual deaths attributable to raised BP", "~10 million", "Leading single modifiable cause of death"],
        ],
        widths=[6.0, 3.5, 6.5],
    )
    caption = document.add_paragraph(
        "Table 2. Headline global indicators for hypertension (WHO Global Report on "
        "Hypertension and related NCD Risk Factor Collaboration analyses).",
        style="Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("4. Pathophysiology and Risk Factors", level=1)
    document.add_paragraph(
        "Arterial pressure is the product of cardiac output and systemic vascular resistance. "
        "In about 90-95 per cent of cases no single cause is identified; this primary "
        "(essential) hypertension results from the interaction of genetic susceptibility with "
        "environmental exposures acting through several overlapping mechanisms: impaired renal "
        "sodium handling and volume expansion, inappropriate activation of the "
        "renin-angiotensin-aldosterone system, sustained sympathetic over-activity, endothelial "
        "dysfunction with reduced nitric-oxide availability, and progressive stiffening of the "
        "large arteries with age."
    )
    add_figure(
        document,
        "fig3_pathophysiology.png",
        "Figure 3. Physiological determinants of blood pressure and the mechanisms that raise "
        "it in primary hypertension.",
    )

    document.add_heading("4.1 Modifiable risk factors", level=2)
    for item in [
        "High dietary sodium intake (most populations consume roughly twice the WHO limit of 5 g of salt per day) and low potassium intake.",
        "Overweight and obesity, especially central adiposity, and insulin resistance.",
        "Physical inactivity and low cardiorespiratory fitness.",
        "Excess alcohol consumption; the 2025 ACC/AHA guideline advises avoidance rather than mere moderation.",
        "Tobacco use, which raises pressure acutely and accelerates arterial damage.",
        "Chronic psychosocial stress, short or poor-quality sleep, and untreated obstructive sleep apnoea.",
    ]:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("4.2 Non-modifiable and secondary contributors", level=2)
    document.add_paragraph(
        "Age, male sex before mid-life, family history, and ancestry-related susceptibility "
        "cannot be changed but sharpen screening priorities. Secondary hypertension accounts "
        "for 5-10 per cent of cases and should be considered when hypertension is severe, "
        "resistant, of abrupt onset, or presents before the age of 30. Recognised causes include "
        "primary aldosteronism (the most common and frequently under-diagnosed), chronic kidney "
        "disease, renal artery stenosis, phaeochromocytoma, Cushing syndrome, thyroid and "
        "parathyroid disorders, coarctation of the aorta, obstructive sleep apnoea, and drugs "
        "such as non-steroidal anti-inflammatory agents, combined oral contraceptives, "
        "decongestants, glucocorticoids, ciclosporin and stimulants."
    )

    document.add_heading("5. Clinical Presentation and Complications", level=1)
    document.add_paragraph(
        "Hypertension is usually asymptomatic, which is why it is described as a silent killer. "
        "When symptoms occur they are non-specific: headache, dizziness, palpitations, epistaxis "
        "or visual disturbance, and they generally reflect severe elevation rather than typical "
        "disease. Consequently the diagnosis depends on measurement, not on how the patient "
        "feels, and the clinical importance of hypertension lies in the cumulative damage it "
        "inflicts on target organs."
    )
    add_figure(
        document,
        "fig4_organ_damage.png",
        "Figure 4. Hypertension-mediated organ damage. Injury develops silently over years "
        "before the first clinical event.",
    )
    document.add_paragraph(
        "Left ventricular hypertrophy, heart failure with preserved or reduced ejection "
        "fraction, atrial fibrillation, ischaemic and haemorrhagic stroke, vascular cognitive "
        "impairment, chronic kidney disease, hypertensive retinopathy, aortic aneurysm and "
        "peripheral artery disease are all strongly and causally linked to long-standing raised "
        "pressure. Early detection of organ damage, for example albuminuria or "
        "electrocardiographic hypertrophy, identifies patients who benefit most from intensive "
        "treatment."
    )

    document.add_heading("6. Diagnosis and Blood-Pressure Measurement", level=1)
    document.add_paragraph(
        "Because treatment decisions rest on the number obtained, measurement technique is a "
        "clinical intervention in its own right. Readings should be taken with a validated "
        "device and correctly sized cuff after five minutes of rest, with the patient seated, "
        "back supported, legs uncrossed, feet flat and the arm supported at heart level, and "
        "without talking. At least two readings should be averaged on two or more occasions."
    )
    document.add_paragraph(
        "Out-of-office measurement is now recommended to confirm the diagnosis. Home monitoring "
        "and 24-hour ambulatory monitoring identify white-coat hypertension, in which office "
        "readings are raised but out-of-office readings are normal, and masked hypertension, in "
        "which the reverse is true and cardiovascular risk is substantially underestimated. "
        "Ambulatory monitoring also reveals nocturnal non-dipping, an independent risk marker."
    )

    add_table(
        document,
        ["Assessment method", "Hypertension threshold", "Main clinical value"],
        [
            ["Standardised office reading", "≥ 140/90 (ESC) or ≥ 130/80 (ACC/AHA)", "Screening and initial classification"],
            ["Home monitoring (average)", "≥ 135/85", "Confirmation, self-management, adherence support"],
            ["Ambulatory, 24-hour average", "≥ 130/80", "Reference standard; detects nocturnal patterns"],
            ["Ambulatory, daytime average", "≥ 135/85", "Confirms white-coat or masked hypertension"],
            ["Ambulatory, night-time average", "≥ 120/70", "Identifies non-dipping and high-risk phenotypes"],
        ],
        widths=[5.0, 5.5, 5.5],
    )
    caption = document.add_paragraph(
        "Table 3. Diagnostic thresholds by measurement method.", style="Caption"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "The baseline work-up complements measurement: history and examination, "
        "electrocardiogram, serum electrolytes, creatinine with estimated glomerular filtration "
        "rate, fasting glucose or HbA1c, lipid profile, urinalysis with urine "
        "albumin-to-creatinine ratio, and a formal cardiovascular risk estimate (PREVENT in the "
        "United States, SCORE2 or SCORE2-OP in Europe). Screening for primary aldosteronism with "
        "an aldosterone-to-renin ratio is recommended in resistant and other high-suspicion "
        "cases."
    )

    document.add_heading("7. Non-Pharmacological Management", level=1)
    document.add_paragraph(
        "Lifestyle therapy is recommended for every person with elevated blood pressure, both as "
        "initial treatment and alongside drugs. Its effects are additive and, taken together, "
        "can match the effect of a single antihypertensive agent."
    )
    add_table(
        document,
        ["Intervention", "Practical target", "Typical systolic reduction"],
        [
            ["Sodium reduction", "< 2 g sodium (~5 g salt) per day", "4-6 mm Hg"],
            ["DASH-type dietary pattern", "Vegetables, fruit, whole grains, low-fat dairy, low saturated fat", "5-11 mm Hg"],
            ["Weight loss", "About 1 mm Hg per kilogram lost", "5 mm Hg or more"],
            ["Aerobic physical activity", "90-150 minutes per week, most days", "5-8 mm Hg"],
            ["Dynamic or isometric resistance training", "2-3 sessions per week", "4-5 mm Hg"],
            ["Alcohol reduction", "Avoid or minimise intake", "4 mm Hg"],
            ["Increased dietary potassium", "3.5-5 g per day unless kidney disease", "4-5 mm Hg"],
        ],
        widths=[4.6, 6.4, 5.0],
    )
    caption = document.add_paragraph(
        "Table 4. Evidence-based lifestyle interventions and their approximate effect on "
        "systolic blood pressure.",
        style="Caption",
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("8. Pharmacological Management", level=1)
    document.add_paragraph(
        "Drug treatment is indicated for all adults with stage 2 hypertension, for adults with "
        "stage 1 hypertension and elevated cardiovascular risk, established cardiovascular "
        "disease, diabetes or chronic kidney disease, and for anyone whose pressure remains "
        "above target after a trial of lifestyle therapy. Contemporary guidelines favour "
        "starting with a low-dose two-drug single-pill combination rather than sequential "
        "monotherapy, because combination therapy lowers pressure faster, is better tolerated "
        "than high-dose monotherapy, and markedly improves adherence."
    )
    add_table(
        document,
        ["Drug class", "Representative agents", "Notes and cautions"],
        [
            ["ACE inhibitors", "Ramipril, lisinopril, perindopril", "Kidney and heart protective; cough; avoid in pregnancy"],
            ["Angiotensin receptor blockers", "Losartan, valsartan, telmisartan", "Similar benefit without cough; avoid in pregnancy"],
            ["Calcium channel blockers", "Amlodipine, nifedipine (long-acting)", "Effective in older and Black patients; ankle oedema"],
            ["Thiazide-type diuretics", "Indapamide, chlortalidone, hydrochlorothiazide", "Monitor sodium, potassium, uric acid and glucose"],
            ["Mineralocorticoid antagonists", "Spironolactone, eplerenone", "Preferred fourth agent in resistant hypertension"],
            ["Beta blockers", "Bisoprolol, metoprolol succinate", "Reserved for specific indications such as ischaemia or heart failure"],
        ],
        widths=[4.4, 5.6, 6.0],
    )
    caption = document.add_paragraph(
        "Table 5. Principal antihypertensive drug classes.", style="Caption"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(
        "The usual treatment goal is a systolic pressure below 130 mm Hg under the ACC/AHA "
        "framework and a target range of 120-129 mm Hg under the ESC framework, provided the "
        "regimen is tolerated; less intensive, individualised goals apply to frail older adults, "
        "people with symptomatic orthostatic hypotension and those with limited life expectancy. "
        "Response should be reviewed after four to eight weeks and treatment escalated promptly, "
        "since therapeutic inertia is a principal reason for poor control."
    )
    add_figure(
        document,
        "fig5_management_pathway.png",
        "Figure 5. A stepped-care pathway from confirmation of the diagnosis to sustained "
        "long-term control.",
    )
    add_figure(
        document,
        "fig6_risk_reduction.png",
        "Figure 6. Approximate relative risk reductions associated with a sustained 10 mm Hg "
        "fall in systolic blood pressure, as reported in large randomised-trial meta-analyses.",
    )

    document.add_heading("9. Special Populations and Resistant Hypertension", level=1)
    document.add_paragraph(
        "Chronic kidney disease requires a renin-angiotensin system blocker when albuminuria is "
        "present, with careful monitoring of potassium and kidney function. In diabetes, blood "
        "pressure lowering delivers cardiovascular and renal benefit that complements glycaemic "
        "control. Older and frail adults benefit from treatment but need slower titration, "
        "standing blood-pressure checks and review of falls risk. In pregnancy, labetalol, "
        "nifedipine and methyldopa are the agents of choice, ACE inhibitors and angiotensin "
        "receptor blockers are contraindicated, and low-dose aspirin is used to reduce the risk "
        "of pre-eclampsia in high-risk women."
    )
    document.add_paragraph(
        "Resistant hypertension is defined as blood pressure that remains above target despite "
        "three appropriately dosed agents including a diuretic, or that requires four or more "
        "agents. Management begins by excluding pseudo-resistance from poor technique, white-coat "
        "effect and non-adherence, then removing interfering substances, screening for secondary "
        "causes, optimising diuretic therapy, and adding spironolactone. Renal denervation is "
        "recognised in recent guidelines as an adjunctive option in carefully selected patients "
        "treated at experienced centres."
    )

    document.add_heading("10. Hypertensive Crises", level=1)
    document.add_paragraph(
        "A blood pressure above 180/120 mm Hg constitutes a crisis. In a hypertensive emergency "
        "there is acute hypertension-mediated organ damage such as encephalopathy, stroke, acute "
        "coronary syndrome, pulmonary oedema, aortic dissection, acute kidney injury or "
        "eclampsia; this requires admission and intravenous therapy with controlled, gradual "
        "reduction, typically no more than 25 per cent within the first hour, except in aortic "
        "dissection or where thrombolysis is planned, when faster and lower targets apply. In "
        "hypertensive urgency there is no acute organ damage, and pressure should be lowered over "
        "hours to days with oral agents and prompt outpatient follow-up; rapid reduction is "
        "harmful and should be avoided."
    )

    document.add_heading("11. Prevention and Public-Health Strategy", level=1)
    document.add_paragraph(
        "Because risk rises continuously, population-wide measures deliver benefits that "
        "individual clinical care cannot. WHO's HEARTS technical package operationalises this "
        "through healthy-lifestyle counselling, evidence-based standard treatment protocols, "
        "access to essential medicines and technology, risk-based management, team-based care "
        "with task sharing, and systems for monitoring. Mandatory sodium reformulation of "
        "processed foods, front-of-pack labelling, tobacco and alcohol control, urban design that "
        "supports physical activity, and inclusion of fixed-dose combinations in essential "
        "medicine lists are among the most cost-effective interventions available to any health "
        "system."
    )
    document.add_paragraph(
        "At the service level, the practical priorities are simple: measure blood pressure "
        "accurately at every opportunity, use a single standard protocol so that any team member "
        "can escalate treatment, prescribe single-pill combinations to reduce pill burden, and "
        "track control rates as a routine quality indicator rather than an audit exercise."
    )

    document.add_heading("12. Conclusion", level=1)
    document.add_paragraph(
        "Hypertension remains the leading modifiable cause of cardiovascular disease and "
        "premature death, yet it is one of the most tractable problems in medicine. The tools "
        "required, accurate measurement, a handful of inexpensive generic drug classes and "
        "well-established lifestyle interventions, are already available almost everywhere. The "
        "decisive gap is implementation: finding people who do not know they are affected, "
        "starting effective combination therapy without delay, titrating to target and sustaining "
        "control over decades. Closing that gap, supported by population-level prevention, would "
        "prevent tens of millions of deaths within a generation."
    )

    document.add_heading("13. References", level=1)
    references = [
        "World Health Organization. Hypertension fact sheet. Geneva: WHO.",
        "World Health Organization. Global Report on Hypertension: the race against a silent killer. Geneva: WHO, 2023.",
        "World Health Organization. HEARTS technical package for cardiovascular disease management in primary health care. Geneva: WHO.",
        "American Heart Association / American College of Cardiology. 2025 Guideline for the Prevention, Detection, Evaluation and Management of High Blood Pressure in Adults. Circulation / Journal of the American College of Cardiology, 2025.",
        "Whelton PK, Carey RM, Aronow WS, et al. 2017 ACC/AHA/AAPA/ABC/ACPM/AGS/APhA/ASH/ASPC/NMA/PCNA Guideline for the Prevention, Detection, Evaluation and Management of High Blood Pressure in Adults. Hypertension. 2018;71(6):e13-e115.",
        "McEvoy JW, McCarthy CP, Bruno RM, et al. 2024 ESC Guidelines for the management of elevated blood pressure and hypertension. European Heart Journal. 2024;45(38):3912-4018.",
        "NCD Risk Factor Collaboration. Worldwide trends in hypertension prevalence and progress in treatment and control from 1990 to 2019. The Lancet. 2021;398(10304):957-980.",
        "Ettehad D, Emdin CA, Kiran A, et al. Blood pressure lowering for prevention of cardiovascular disease and death: a systematic review and meta-analysis. The Lancet. 2016;387(10022):957-967.",
        "Blood Pressure Lowering Treatment Trialists' Collaboration. Pharmacological blood pressure lowering for primary and secondary prevention of cardiovascular disease across different levels of blood pressure. The Lancet. 2021;397(10285):1625-1636.",
        "SPRINT Research Group. A randomized trial of intensive versus standard blood-pressure control. New England Journal of Medicine. 2015;373(22):2103-2116.",
        "Appel LJ, Moore TJ, Obarzanek E, et al. A clinical trial of the effects of dietary patterns on blood pressure (DASH). New England Journal of Medicine. 1997;336(16):1117-1124.",
        "Khan SS, Matsushita K, Sang Y, et al. Development and validation of the American Heart Association PREVENT equations. Circulation. 2024;149(6):430-449.",
    ]
    for index, reference in enumerate(references, start=1):
        paragraph = document.add_paragraph(f"[{index}] {reference}")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.runs[0].font.size = Pt(9.5)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    run = note.add_run(
        "All figures in this report were produced by the authors specifically for this "
        "document; no third-party images are reproduced. This report is an educational "
        "evidence review and does not replace individual clinical judgement."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def build_docx() -> Path:
    document = Document()
    properties = document.core_properties
    properties.title = "Hypertension: Epidemiology, Pathophysiology, Diagnosis and Evidence-Based Management"
    properties.subject = "Clinical research report on hypertension"
    properties.category = "Clinical research report"
    properties.keywords = "hypertension; blood pressure; cardiovascular risk; WHO; ACC/AHA; ESC"
    configure_styles(document)
    section = configure_page(document)
    add_header_footer(section, "Hypertension - Clinical Research Report")
    add_cover(document)
    add_table_of_contents(document)
    add_body(document)
    document.save(DOCX)
    print(f"wrote {DOCX}")
    return DOCX


def convert_to_pdf(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) is required to export the PDF")
    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(docx_path.parent),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
    )
    pdf_path = docx_path.with_suffix(".pdf")
    print(f"wrote {pdf_path}")
    return pdf_path


if __name__ == "__main__":
    convert_to_pdf(build_docx())
